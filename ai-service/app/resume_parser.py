
from io import BytesIO

import pymupdf
import pytesseract
from docx import Document
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader


MAX_PDF_PAGES_FOR_OCR = 10


def extract_resume_text(content: bytes, extension: str) -> str:
    """Extract text from PDF, DOCX, JPG, or JPEG resumes."""

    extension = extension.lower()

    if extension == ".docx":
        document = Document(BytesIO(content))
        parts = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        text = "\n".join(parts)

    elif extension == ".pdf":
        reader = PdfReader(BytesIO(content))

        if reader.is_encrypted:
            raise ValueError("Password-protected PDFs are not supported.")

        text = "\n".join(
            page.extract_text() or "" for page in reader.pages
        ).strip()

        # Use OCR when the PDF has no extractable text.
        if not text:
            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                if len(pdf) > MAX_PDF_PAGES_FOR_OCR:
                    raise ValueError(
                        "Scanned PDFs are limited to 10 pages for OCR."
                    )

                pages = []

                for page in pdf:
                    pixmap = page.get_pixmap(
                        matrix=pymupdf.Matrix(1.5, 1.5),
                        alpha=False,
                    )
                    image = Image.open(
                        BytesIO(pixmap.tobytes("png"))
                    )
                    pages.append(pytesseract.image_to_string(image))

                text = "\n".join(pages)

    elif extension in {".jpg", ".jpeg"}:
        try:
            with Image.open(BytesIO(content)) as image:
                image.verify()

            with Image.open(BytesIO(content)) as image:
                text = pytesseract.image_to_string(image)

        except (UnidentifiedImageError, OSError) as exc:
            raise ValueError(
                "The uploaded image could not be read."
            ) from exc

    else:
        raise ValueError("Unsupported file format.")

    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text found. Try a clearer image or PDF."
        )

    return text
